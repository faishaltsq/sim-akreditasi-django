import json
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.core.exceptions import PermissionDenied
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.db import transaction
from django.views.decorators.http import require_POST
from django.db.models import Count, Q, Sum
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import (
    Category, StandardItem, QualityRecord,
    EvidenceReq, EvidenceFile, UnitKerja, Framework, AuditLog, RumahSakitProfile
)
from .forms import StandardItemForm, QualityRecordForm, UnitKerjaForm, EvidenceFileUploadForm
from .supabase_storage import upload_to_supabase_storage


# ==============================================================================
# RBAC HELPERS
# ==============================================================================
def _is_admin(user):
    """Admin RS atau Superuser."""
    if not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    p = getattr(user, 'profile', None)
    return p and p.role in ('SUPER_ADMIN', 'ADMIN_RS')


def _can_edit(user):
    """Role yang boleh create/edit/delete EP (bukan Asesor, bukan Nakes)."""
    if not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    p = getattr(user, 'profile', None)
    return p and p.can_edit


def _is_nakes_or_admin(user):
    """Nakes sendiri atau Admin."""
    if not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    p = getattr(user, 'profile', None)
    return p and p.role in ('SUPER_ADMIN', 'ADMIN_RS', 'STAF_NAKES')


admin_required = user_passes_test(_is_admin, login_url='/accounts/login/')
editor_required = user_passes_test(_can_edit, login_url='/accounts/login/')
nakes_or_admin = user_passes_test(_is_nakes_or_admin, login_url='/accounts/login/')


# ==============================================================================
# 1. DASHBOARD UTAMA (SIK AP Overview)
# ==============================================================================
@login_required
def dashboard(request):
    # Smart redirect: nakes langsung ke portal unit
    profile = getattr(request.user, 'profile', None)
    if profile and profile.role == 'STAF_NAKES':
        return redirect('akreditasi:portal_nakes')

    rs_profile = RumahSakitProfile.get_default()
    framework = Framework.objects.first()
    categories = Category.objects.all().order_by('order')

    all_items = StandardItem.objects.select_related('category', 'record__unit').prefetch_related('evidence_reqs__files')
    total_ep = all_items.count()

    total_possible_score = total_ep * 10
    total_current_score = sum(getattr(i, 'record', None).score if hasattr(i, 'record') and i.record else 0 for i in all_items)
    capaian_rata_rata = round((total_current_score / total_possible_score) * 100, 1) if total_possible_score > 0 else 0

    # Total Dokumen: file bukti terunggah vs target (total evidence requirements)
    all_reqs = EvidenceReq.objects.all()
    total_target_dokumen = all_reqs.count()
    total_uploaded_dokumen = EvidenceFile.objects.values('requirement_id').distinct().count()

    # Evaluasi PDCA: % EP yang sudah masuk tahap Check atau Action
    ep_check_action = QualityRecord.objects.filter(
        Q(eval_notes__isnull=False, eval_notes__gt='') |
        Q(action_plan__isnull=False, action_plan__gt='') |
        Q(score__gt=0)
    ).count()
    persen_evaluasi_pdca = round((ep_check_action / total_ep) * 100, 1) if total_ep > 0 else 0

    # Anggaran RKA
    total_anggaran = QualityRecord.objects.aggregate(total=Sum('est_cost'))['total'] or 0
    anggaran_capex = QualityRecord.objects.filter(
        Q(budget_source__icontains='Capex') | Q(budget_source__icontains='Sarpras')
    ).aggregate(total=Sum('est_cost'))['total'] or 0
    anggaran_opex = max(0, float(total_anggaran) - float(anggaran_capex))
    non_biaya_count = QualityRecord.objects.filter(est_cost=0).count()

    # Data Pie Chart "Status Kepatuhan EP"
    count_tercapai = QualityRecord.objects.filter(score=10).count()
    count_proses = QualityRecord.objects.filter(score=5).count()
    count_belum = total_ep - (count_tercapai + count_proses)
    if total_ep > 0:
        pct_tercapai = round((count_tercapai / total_ep) * 100, 1)
        pct_proses = round((count_proses / total_ep) * 100, 1)
        pct_belum = round((count_belum / total_ep) * 100, 1)
    else:
        pct_tercapai = pct_proses = pct_belum = 0

    # Aktivitas Terkini (5 log audit terbaru)
    aktivitas_terkini = AuditLog.objects.select_related('user').order_by('-timestamp')[:5]

    # Struktur Unit Kerja RS untuk Dashboard (Pohon Hierarki)
    root_units = UnitKerja.objects.filter(level=1).prefetch_related('children__children').order_by('code')
    total_units_count = UnitKerja.objects.count()

    context = {
        'rs_profile': rs_profile,
        'framework': framework,
        'categories': categories,
        'total_ep': total_ep,
        'capaian_rata_rata': capaian_rata_rata,
        'total_uploaded_dokumen': total_uploaded_dokumen,
        'total_target_dokumen': total_target_dokumen,
        'persen_evaluasi_pdca': persen_evaluasi_pdca,
        'ep_check_action': ep_check_action,
        'total_anggaran': float(total_anggaran),
        'anggaran_capex': float(anggaran_capex),
        'anggaran_opex': float(anggaran_opex),
        'non_biaya_count': non_biaya_count,
        'count_tercapai': count_tercapai,
        'count_proses': count_proses,
        'count_belum': count_belum,
        'pct_tercapai': pct_tercapai,
        'pct_proses': pct_proses,
        'pct_belum': pct_belum,
        'aktivitas_terkini': aktivitas_terkini,
        'root_units': root_units,
        'total_units_count': total_units_count,
    }
    return render(request, 'akreditasi/dashboard.html', context)


# ==============================================================================
# 2. HALAMAN INTI: MATRIKS PDCA (Detail Per Pokja & Standar)
# ==============================================================================
@login_required
def matriks_pdca(request):
    categories = Category.objects.all().order_by('order')
    if not categories.exists():
        return render(request, 'akreditasi/empty.html')

    selected_cat_id = request.GET.get('cat')
    selected_sub = request.GET.get('sub', 'ALL')
    selected_unit_id = request.GET.get('unit', 'ALL')

    if selected_cat_id:
        active_category = get_object_or_404(Category, id=selected_cat_id)
    else:
        active_category = categories.first()

    # Ambil list sub-standar unik di Pokja ini untuk filter
    sub_standards = (
        StandardItem.objects.filter(category=active_category)
        .values_list('sub_standard', flat=True)
        .distinct()
        .order_by('sub_standard')
    )

    items_qs = active_category.items.prefetch_related(
        'evidence_reqs__files',
        'record__unit'
    ).order_by('order', 'code')

    if selected_sub != 'ALL':
        items_qs = items_qs.filter(sub_standard=selected_sub)

    if selected_unit_id != 'ALL':
        items_qs = items_qs.filter(record__unit_id=selected_unit_id)

    items = list(items_qs)

    # Kalkulasi metriks pokja aktif
    total_possible = len(items) * 10
    total_score = sum(getattr(item, 'record', None).score if hasattr(item, 'record') and item.record else 0 for item in items)
    percentage = round((total_score / total_possible) * 100) if total_possible > 0 else 0

    units = UnitKerja.objects.all().order_by('code')

    context = {
        'categories': categories,
        'active_category': active_category,
        'sub_standards': sub_standards,
        'selected_sub': selected_sub,
        'selected_unit_id': selected_unit_id,
        'units': units,
        'items': items,
        'total_possible': total_possible,
        'total_score': total_score,
        'percentage': percentage,
    }
    return render(request, 'akreditasi/matriks_pdca.html', context)


# ==============================================================================
# 3. QUICK SCORE AJAX (Dropdown Skor 10, 5, 0 Langsung Simpan)
# ==============================================================================
@login_required
@require_POST
def quick_score(request, record_id):
    record = get_object_or_404(QualityRecord, id=record_id)
    try:
        data = json.loads(request.body)
        new_score = int(data.get('score', 0))
        if new_score in [10, 5, 0]:
            old_score = record.score
            record.score = new_score
            record.save()

            AuditLog.objects.create(
                user=request.user,
                aksi='SCORE',
                model_name='QualityRecord',
                object_repr=f"{record.standard_item.code} - {record.unit.name}",
                detail=f"Skor diubah dari {old_score} ke {new_score}"
            )

            cat = record.standard_item.category
            return JsonResponse({
                'success': True,
                'score': new_score,
                'cat_score': cat.total_score,
                'cat_max': cat.max_possible_score,
                'cat_percentage': cat.percentage,
            })
        return JsonResponse({'success': False, 'error': 'Invalid score'}, status=400)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


# ==============================================================================
# 4. INLINE EDIT SEL PDCA (Plan, Do, Check, Action AJAX)
# ==============================================================================
@login_required
@require_POST
def inline_edit_pdca(request, record_id):
    record = get_object_or_404(QualityRecord, id=record_id)
    try:
        data = json.loads(request.body)
        field = data.get('field')  # 'baseline_data', 'quality_target', 'risk_mitigation', 'eval_notes', 'action_plan'
        value = data.get('value', '').strip()

        allowed_fields = ['baseline_data', 'quality_target', 'risk_mitigation', 'eval_notes', 'action_plan', 'pic', 'target_date']
        if field in allowed_fields:
            setattr(record, field, value)
            record.save()

            AuditLog.objects.create(
                user=request.user,
                aksi='UPDATE',
                model_name='QualityRecord',
                object_repr=f"{record.standard_item.code} ({field})",
                detail=f"Update isi sel {field}: {value[:80]}"
            )
            return JsonResponse({'success': True, 'field': field, 'value': value})
        return JsonResponse({'success': False, 'error': 'Field tidak diizinkan'}, status=400)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


# ==============================================================================
# 5. MODAL MANAJEMEN BUKTI DOKUMEN (3-TAB: Upload Baru, Pilih Repo, Terlink)
# ==============================================================================
@login_required
def modal_bukti_data(request, item_id):
    item = get_object_or_404(StandardItem, id=item_id)
    reqs = item.evidence_reqs.prefetch_related('files').all()

    # Tab 3: Daftar berkas terlink ke EP ini
    linked_files = []
    for req in reqs:
        for f in req.files.all():
            linked_files.append({
                'id': f.id,
                'req_id': req.id,
                'req_type': req.category_type,
                'req_title': req.title,
                'file_name': f.file_name,
                'file_size': f.file_size,
                'file_url': f.get_download_url,
                'version': f.version,
                'status': f.status,
                'uploaded_at': f.uploaded_at.strftime('%d/%m/%Y %H:%M'),
            })

    # Tab 2: Dokumen repositori lain yang sudah ada di sistem (bisa ditautkan)
    all_repo_files = EvidenceFile.objects.select_related('requirement__standard_item').all().order_by('-uploaded_at')[:30]
    repo_list = []
    current_linked_ids = [f['id'] for f in linked_files]
    for rf in all_repo_files:
        if rf.id not in current_linked_ids:
            repo_list.append({
                'id': rf.id,
                'file_name': rf.file_name,
                'file_size': rf.file_size,
                'ep_origin': rf.requirement.standard_item.code,
                'type': rf.requirement.category_type,
            })

    req_options = [{'id': r.id, 'type': r.category_type, 'title': r.title} for r in reqs]

    return JsonResponse({
        'item_code': item.code,
        'item_desc': item.description,
        'reqs': req_options,
        'linked_files': linked_files,
        'repo_files': repo_list,
    })


@login_required
@require_POST
def upload_bukti_ajax(request, req_id):
    evidence_req = get_object_or_404(EvidenceReq, id=req_id)
    file_obj = request.FILES.get('file')
    doc_name = request.POST.get('doc_name', '').strip()

    if not file_obj:
        return JsonResponse({'success': False, 'error': 'Pilih berkas terlebih dahulu.'}, status=400)

    final_name = doc_name if doc_name else file_obj.name
    res = upload_to_supabase_storage(file_obj, final_name)
    if res.get('success'):
        size_mb = f"{file_obj.size / (1024 * 1024):.1f} MB" if file_obj.size >= 1024 * 1024 else f"{file_obj.size / 1024:.0f} KB"
        ev_file = EvidenceFile.objects.create(
            requirement=evidence_req,
            file=file_obj if not res.get('is_cloud') else None,
            file_url=res.get('url', ''),
            file_name=final_name,
            file_size=size_mb,
            status='VALID'
        )
        AuditLog.objects.create(
            user=request.user,
            aksi='UPLOAD',
            model_name='EvidenceFile',
            object_repr=f"{evidence_req.standard_item.code} - {final_name}",
            detail="Berkas bukti diunggah via Modal Manajemen Bukti"
        )
        return JsonResponse({
            'success': True,
            'file_id': ev_file.id,
            'file_name': ev_file.file_name,
            'file_url': ev_file.get_download_url,
            'file_size': ev_file.file_size,
        })
    return JsonResponse({'success': False, 'error': res.get('error', 'Gagal upload')}, status=500)


@login_required
@require_POST
def link_existing_doc(request):
    try:
        data = json.loads(request.body)
        target_req_id = data.get('req_id')
        source_file_id = data.get('file_id')

        target_req = get_object_or_404(EvidenceReq, id=target_req_id)
        source_file = get_object_or_404(EvidenceFile, id=source_file_id)

        # Clone record EvidenceFile ke target requirement
        new_file = EvidenceFile.objects.create(
            requirement=target_req,
            file=source_file.file,
            file_url=source_file.file_url,
            file_name=source_file.file_name,
            file_size=source_file.file_size,
            version=source_file.version,
            status=source_file.status
        )

        AuditLog.objects.create(
            user=request.user,
            aksi='UPLOAD',
            model_name='EvidenceFile',
            object_repr=f"{target_req.standard_item.code} - {new_file.file_name}",
            detail=f"Menautkan dokumen dari repositori ({source_file.file_name})"
        )
        return JsonResponse({'success': True, 'file_id': new_file.id, 'file_name': new_file.file_name})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_POST
def unlink_doc(request, file_id):
    ev_file = get_object_or_404(EvidenceFile, id=file_id)
    name = ev_file.file_name
    ep_code = ev_file.requirement.standard_item.code
    ev_file.delete()

    AuditLog.objects.create(
        user=request.user,
        aksi='DELETE',
        model_name='EvidenceFile',
        object_repr=f"{ep_code} - {name}",
        detail="Membatalkan tautan dokumen bukti"
    )
    return JsonResponse({'success': True, 'message': f"Dokumen {name} berhasil dilepas dari tautan."})


# ==============================================================================
# 6. DAFTAR ELEMEN PENILAIAN (EP LIST VIEW)
# ==============================================================================
@login_required
def ep_list(request):
    categories = Category.objects.all().order_by('order')
    units = UnitKerja.objects.all().order_by('code')

    selected_cat = request.GET.get('cat', 'ALL')
    selected_sub = request.GET.get('sub', 'ALL')
    selected_unit = request.GET.get('unit', 'ALL')
    search_q = request.GET.get('q', '').strip()

    items_qs = StandardItem.objects.select_related('category', 'record__unit').prefetch_related('evidence_reqs__files').order_by('category__order', 'order', 'code')

    if selected_cat != 'ALL':
        items_qs = items_qs.filter(category_id=selected_cat)
    if selected_sub != 'ALL':
        items_qs = items_qs.filter(sub_standard=selected_sub)
    if selected_unit != 'ALL':
        items_qs = items_qs.filter(record__unit_id=selected_unit)
    if search_q:
        items_qs = items_qs.filter(
            Q(code__icontains=search_q) |
            Q(description__icontains=search_q) |
            Q(sub_title__icontains=search_q)
        )

    # Sub-standard list for filtering
    sub_list = StandardItem.objects.values_list('sub_standard', flat=True).distinct().order_by('sub_standard')

    return render(request, 'akreditasi/ep_list.html', {
        'items': items_qs,
        'categories': categories,
        'units': units,
        'sub_list': sub_list,
        'selected_cat': selected_cat,
        'selected_sub': selected_sub,
        'selected_unit': selected_unit,
        'search_q': search_q,
    })


# ==============================================================================
# 7. REPOSITORI DOKUMEN BUKTI (RDWOS HUB)
# ==============================================================================
@login_required
def dokumen_hub(request):
    selected_type = request.GET.get('type', 'ALL')
    search_q = request.GET.get('q', '').strip()

    files_qs = EvidenceFile.objects.select_related(
        'requirement__standard_item__category',
        'requirement__standard_item__record__unit'
    ).order_by('-uploaded_at')

    if selected_type != 'ALL':
        files_qs = files_qs.filter(requirement__category_type=selected_type)

    if search_q:
        files_qs = files_qs.filter(
            Q(file_name__icontains=search_q) |
            Q(requirement__title__icontains=search_q) |
            Q(requirement__standard_item__code__icontains=search_q)
        )

    return render(request, 'akreditasi/dokumen_hub.html', {
        'files': files_qs,
        'selected_type': selected_type,
        'search_q': search_q,
        'total_files': files_qs.count(),
    })


# ==============================================================================
# 8. FORM CREATE & EDIT EP
# ==============================================================================
@login_required
@editor_required
def ep_create(request):
    if request.method == 'POST':
        item_form = StandardItemForm(request.POST)
        record_form = QualityRecordForm(request.POST)
        if item_form.is_valid() and record_form.is_valid():
            with transaction.atomic():
                item = item_form.save()
                record = record_form.save(commit=False)
                record.standard_item = item
                record.save()

                req_types = request.POST.getlist('req_type[]')
                req_titles = request.POST.getlist('req_title[]')
                for t, title in zip(req_types, req_titles):
                    if t and title:
                        EvidenceReq.objects.create(
                            standard_item=item,
                            category_type=t,
                            title=title,
                            is_mandatory=True
                        )

                AuditLog.objects.create(
                    user=request.user,
                    aksi='CREATE',
                    model_name='StandardItem',
                    object_repr=item.code,
                    detail=f"EP baru ditambahkan: {item.description[:100]}"
                )

                messages.success(request, f"Elemen Penilaian {item.code} berhasil ditambahkan!")
                return redirect(f"/matriks/?cat={item.category.id}")
    else:
        init_cat = request.GET.get('cat')
        initial_data = {}
        if init_cat:
            initial_data['category'] = init_cat
        item_form = StandardItemForm(initial=initial_data)
        record_form = QualityRecordForm()

    return render(request, 'akreditasi/ep_form.html', {
        'item_form': item_form,
        'record_form': record_form,
        'title': 'Tambah Elemen Penilaian Baru',
    })


@login_required
@editor_required
def ep_edit(request, item_id):
    item = get_object_or_404(StandardItem, id=item_id)
    record, _ = QualityRecord.objects.get_or_create(
        standard_item=item,
        defaults={'unit': UnitKerja.objects.first(), 'score': 0}
    )

    if request.method == 'POST':
        item_form = StandardItemForm(request.POST, instance=item)
        record_form = QualityRecordForm(request.POST, instance=record)
        if item_form.is_valid() and record_form.is_valid():
            with transaction.atomic():
                item_form.save()
                record_form.save()

                AuditLog.objects.create(
                    user=request.user,
                    aksi='UPDATE',
                    model_name='StandardItem',
                    object_repr=item.code,
                    detail="Data EP dan Siklus PDCA diperbarui"
                )

                messages.success(request, f"Perubahan pada {item.code} berhasil disimpan!")
                return redirect(f"/matriks/?cat={item.category.id}")
    else:
        item_form = StandardItemForm(instance=item)
        record_form = QualityRecordForm(instance=record)

    return render(request, 'akreditasi/ep_form.html', {
        'item_form': item_form,
        'record_form': record_form,
        'item': item,
        'title': f'Edit EP: {item.code}',
    })


@login_required
def ep_detail(request, item_id):
    item = get_object_or_404(StandardItem, id=item_id)
    record = getattr(item, 'record', None)
    upload_form = EvidenceFileUploadForm()
    return render(request, 'akreditasi/ep_detail.html', {
        'item': item,
        'record': record,
        'upload_form': upload_form,
    })


@login_required
@require_POST
def upload_bukti(request, req_id):
    evidence_req = get_object_or_404(EvidenceReq, id=req_id)
    file_obj = request.FILES.get('file')

    if not file_obj:
        messages.error(request, "Pilih berkas terlebih dahulu.")
        return redirect('akreditasi:ep_detail', item_id=evidence_req.standard_item.id)

    res = upload_to_supabase_storage(file_obj, file_obj.name)
    if res.get('success'):
        size_mb = f"{file_obj.size / (1024 * 1024):.1f} MB" if file_obj.size >= 1024 * 1024 else f"{file_obj.size / 1024:.0f} KB"
        EvidenceFile.objects.create(
            requirement=evidence_req,
            file=file_obj if not res.get('is_cloud') else None,
            file_url=res.get('url', ''),
            file_name=file_obj.name,
            file_size=size_mb,
            status='VALID'
        )
        AuditLog.objects.create(
            user=request.user,
            aksi='UPLOAD',
            model_name='EvidenceFile',
            object_repr=f"{evidence_req.standard_item.code} - {file_obj.name}",
            detail="Berkas bukti berhasil diunggah"
        )
        messages.success(request, f"Berkas '{file_obj.name}' berhasil diunggah!")
    else:
        messages.error(request, f"Gagal mengunggah berkas: {res.get('error', 'Unknown error')}")

    return redirect('akreditasi:ep_detail', item_id=evidence_req.standard_item.id)


@login_required
@editor_required
@require_POST
def ep_delete(request, item_id):
    item = get_object_or_404(StandardItem, id=item_id)
    code = item.code
    cat_id = item.category.id
    item.delete()

    AuditLog.objects.create(
        user=request.user,
        aksi='DELETE',
        model_name='StandardItem',
        object_repr=code,
        detail="Elemen Penilaian dihapus permanen"
    )
    messages.success(request, f"Elemen Penilaian {code} telah dihapus.")
    return redirect(f"/matriks/?cat={cat_id}")


# ==============================================================================
# 9. MANAJEMEN DATA: UNIT KERJA HIERARKI & POKJA
# ==============================================================================
@login_required
@admin_required
def unit_list_create(request):
    if request.method == 'POST':
        form = UnitKerjaForm(request.POST)
        if form.is_valid():
            u = form.save()
            AuditLog.objects.create(
                user=request.user,
                aksi='CREATE',
                model_name='UnitKerja',
                object_repr=f"[{u.code}] {u.name}",
                detail=f"Unit Kerja baru ditambahkan (Level {u.level})"
            )
            messages.success(request, f"Unit Kerja '{u.name}' berhasil ditambahkan!")
            return redirect('akreditasi:unit_list')
    else:
        form = UnitKerjaForm()

    return render(request, 'akreditasi/unit_form.html', {'form': form})


@login_required
def unit_tree(request):
    root_units = UnitKerja.objects.filter(parent__isnull=True).prefetch_related(
        'children__children'
    ).order_by('code')
    all_units = UnitKerja.objects.all().order_by('level', 'code')
    return render(request, 'akreditasi/unit_tree.html', {
        'root_units': root_units,
        'all_units': all_units,
        'total_units': all_units.count(),
    })


@login_required
@admin_required
def unit_edit(request, unit_id):
    unit = get_object_or_404(UnitKerja, id=unit_id)
    if request.method == 'POST':
        form = UnitKerjaForm(request.POST, instance=unit)
        if form.is_valid():
            u = form.save()
            AuditLog.objects.create(
                user=request.user,
                aksi='UPDATE',
                model_name='UnitKerja',
                object_repr=f"[{u.code}] {u.name}",
                detail=f"Data Unit Kerja diperbarui"
            )
            messages.success(request, f"Unit Kerja '{u.name}' berhasil diperbarui!")
            return redirect('akreditasi:unit_tree')
    else:
        form = UnitKerjaForm(instance=unit)

    return render(request, 'akreditasi/unit_form.html', {
        'form': form,
        'unit': unit,
        'title': f'Edit Unit Kerja: {unit.code}',
    })


@login_required
@admin_required
@require_POST
def unit_delete(request, unit_id):
    unit = get_object_or_404(UnitKerja, id=unit_id)
    name = f"[{unit.code}] {unit.name}"
    child_count = unit.children.count()
    if child_count > 0:
        messages.error(request, f"Unit '{name}' masih memiliki {child_count} sub-unit. Hapus sub-unit terlebih dahulu.")
        return redirect('akreditasi:unit_tree')
    if unit.records.exists():
        messages.error(request, f"Unit '{name}' masih terhubung ke Record PDCA. Pindahkan record terlebih dahulu.")
        return redirect('akreditasi:unit_tree')

    unit.delete()
    AuditLog.objects.create(
        user=request.user,
        aksi='DELETE',
        model_name='UnitKerja',
        object_repr=name,
        detail="Unit Kerja dihapus permanen"
    )
    messages.success(request, f"Unit Kerja '{name}' telah dihapus.")
    return redirect('akreditasi:unit_tree')


@login_required
@admin_required
def pokja_manage(request):
    categories = Category.objects.all().order_by('order')
    return render(request, 'akreditasi/pokja_list.html', {'categories': categories})


# ==============================================================================
# 10. PENGATURAN: PROFIL RS & LOG SISTEM
# ==============================================================================
@login_required
@admin_required
def profil_rs_view(request):
    profile = RumahSakitProfile.get_default()
    if request.method == 'POST':
        profile.name = request.POST.get('name', profile.name).strip()
        profile.kode_rs = request.POST.get('kode_rs', profile.kode_rs).strip()
        profile.alamat = request.POST.get('alamat', profile.alamat).strip()
        profile.kota = request.POST.get('kota', profile.kota).strip()
        profile.telepon = request.POST.get('telepon', profile.telepon).strip()
        profile.email = request.POST.get('email', profile.email).strip()
        profile.website = request.POST.get('website', profile.website).strip()
        profile.direktur = request.POST.get('direktur', profile.direktur).strip()
        profile.tipe = request.POST.get('tipe', profile.tipe).strip()
        profile.akreditasi_tahun = int(request.POST.get('akreditasi_tahun', profile.akreditasi_tahun) or 2026)
        profile.logo_url = request.POST.get('logo_url', profile.logo_url).strip()
        profile.save()

        AuditLog.objects.create(
            user=request.user,
            aksi='UPDATE',
            model_name='RumahSakitProfile',
            object_repr=profile.name,
            detail="Profil Rumah Sakit diperbarui"
        )
        messages.success(request, "Profil Rumah Sakit berhasil diperbarui!")
        return redirect('akreditasi:profil_rs')

    return render(request, 'akreditasi/profil_rs.html', {'profile': profile})


@login_required
@admin_required
def audit_log_view(request):
    logs = AuditLog.objects.select_related('user').order_by('-timestamp')[:100]
    return render(request, 'akreditasi/audit_log.html', {'logs': logs})


# ==============================================================================
# 11. LAPORAN & EKSPOR DATA
# ==============================================================================
@login_required
def rekap_view(request):
    categories = Category.objects.all().order_by('order')
    items = StandardItem.objects.select_related('category', 'record__unit').prefetch_related('evidence_reqs__files').order_by('category__order', 'order', 'code')

    total_possible = items.count() * 10
    total_score = sum(i.record.score for i in items if hasattr(i, 'record') and i.record)
    overall_pct = round((total_score / total_possible) * 100) if total_possible > 0 else 0

    total_rka = sum(float(i.record.est_cost or 0) for i in items if hasattr(i, 'record') and i.record)

    return render(request, 'akreditasi/rekap.html', {
        'categories': categories,
        'items': items,
        'total_possible': total_possible,
        'total_score': total_score,
        'overall_pct': overall_pct,
        'total_rka': total_rka,
    })


@login_required
def cetak_dokumen_pokja(request, cat_id):
    """Fase 3.3: Cetak dokumen standar akreditasi per Pokja (PDF / Print layout)."""
    cat = get_object_or_404(Category, id=cat_id)
    items = StandardItem.objects.filter(category=cat).select_related('record__unit').prefetch_related('evidence_reqs__files').order_by('order', 'code')
    ep_count = items.count()
    tercapai = items.filter(record__score=10).count()
    proses = items.filter(record__score=5).count()
    belum = ep_count - tercapai - proses

    return render(request, 'akreditasi/cetak_dokumen.html', {
        'cat': cat,
        'items': items,
        'ep_count': ep_count,
        'tercapai': tercapai,
        'proses': proses,
        'belum': belum,
    })


@login_required
def auto_scoring_pokja(request):
    """Fase 3.4: Auto-scoring pemenuhan EP per Pokja dengan formula KARS STARKES."""
    categories = Category.objects.all().order_by('order')

    def _nilai_akreditasi(pct):
        if pct >= 80: return ('Paripurna', '#16a34a', '★★★★★')
        if pct >= 60: return ('Utama', '#0d9488', '★★★★')
        if pct >= 40: return ('Madya', '#ca8a04', '★★★')
        if pct >= 20: return ('Dasar', '#ea580c', '★★')
        return ('Tidak Terakreditasi', '#dc2626', '★')

    pokja_data = []
    grand_total_possible = 0
    grand_total_score = 0

    for cat in categories:
        cat_items = StandardItem.objects.filter(category=cat).select_related('record')
        ep_count = cat_items.count()
        possible = ep_count * 10
        score = sum(i.record.score for i in cat_items if hasattr(i, 'record') and i.record)
        pct = round((score / possible) * 100, 1) if possible > 0 else 0
        tercapai = cat_items.filter(record__score=10).count()
        proses = cat_items.filter(record__score=5).count()
        belum = ep_count - tercapai - proses
        label, color, stars = _nilai_akreditasi(pct)

        # Hitung kelengkapan bukti per pokja
        req_count = EvidenceReq.objects.filter(standard_item__category=cat).count()
        bukti_count = EvidenceFile.objects.filter(
            requirement__standard_item__category=cat
        ).values('requirement_id').distinct().count()
        bukti_pct = round((bukti_count / req_count) * 100, 1) if req_count > 0 else 0

        grand_total_possible += possible
        grand_total_score += score

        pokja_data.append({
            'cat': cat,
            'ep_count': ep_count,
            'possible': possible,
            'score': score,
            'pct': pct,
            'tercapai': tercapai,
            'proses': proses,
            'belum': belum,
            'label': label,
            'color': color,
            'stars': stars,
            'req_count': req_count,
            'bukti_count': bukti_count,
            'bukti_pct': bukti_pct,
        })

    grand_pct = round((grand_total_score / grand_total_possible) * 100, 1) if grand_total_possible > 0 else 0
    grand_label, grand_color, grand_stars = _nilai_akreditasi(grand_pct)

    return render(request, 'akreditasi/auto_scoring.html', {
        'pokja_data': pokja_data,
        'grand_pct': grand_pct,
        'grand_label': grand_label,
        'grand_color': grand_color,
        'grand_stars': grand_stars,
        'grand_total_score': grand_total_score,
        'grand_total_possible': grand_total_possible,
    })


@login_required
def export_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriks PDCA STARKES"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    headers = [
        "Pokja", "Sub-Standar", "Kode EP", "Unit Kerja", "Asesmen Baseline",
        "PLAN: Indikator Mutu", "PLAN: Mitigasi Risiko",
        "DO: Bukti Pembuktian", "CHECK: Skor (0/5/10)", "CHECK: Catatan",
        "ACTION: RTL", "PIC", "Target Waktu", "Estimasi Biaya (Rp)", "Sumber Anggaran", "Status Anggaran"
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    items = StandardItem.objects.select_related('category', 'record__unit').prefetch_related('evidence_reqs').order_by('category__order', 'order')

    for item in items:
        rec = getattr(item, 'record', None)
        bukti_list = [f"[{req.category_type}] {req.title}" for req in item.evidence_reqs.all()]
        bukti_str = "\n".join(bukti_list)

        row = [
            item.category.code,
            item.sub_standard,
            item.code,
            rec.unit.name if rec and rec.unit else "-",
            rec.baseline_data if rec else "-",
            rec.quality_target if rec else "-",
            rec.risk_mitigation if rec else "-",
            bukti_str,
            rec.score if rec else 0,
            rec.eval_notes if rec else "-",
            rec.action_plan if rec else "-",
            rec.pic if rec else "-",
            rec.target_date if rec else "-",
            float(rec.est_cost or 0) if rec else 0,
            rec.budget_source if rec else "-",
            rec.get_budget_status_display() if rec else "-",
        ]
        ws.append(row)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border_thin
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 25
    ws.column_dimensions['G'].width = 25
    ws.column_dimensions['H'].width = 30
    ws.column_dimensions['I'].width = 12
    ws.column_dimensions['J'].width = 25
    ws.column_dimensions['K'].width = 30
    ws.column_dimensions['L'].width = 15
    ws.column_dimensions['M'].width = 12
    ws.column_dimensions['N'].width = 18
    ws.column_dimensions['O'].width = 20
    ws.column_dimensions['P'].width = 15

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="Matriks_PDCA_Akreditasi_RS.xlsx"'
    wb.save(response)
    return response


# ==============================================================================
# 12. PORTAL NAKES (Capaian Unit + SOP + Upload Bukti + Portofolio KPS)
# ==============================================================================
@login_required
@nakes_or_admin
def portal_nakes(request):
    profile = getattr(request.user, 'profile', None)
    if not profile:
        messages.warning(request, "Profil pengguna belum dibuat.")
        return redirect('akreditasi:dashboard_alias')

    unit = profile.unit_kerja
    rs_profile = RumahSakitProfile.get_default()

    # Data EP unit
    unit_records = []
    unit_items = []
    total_score = 0
    total_possible = 0
    count_10 = count_5 = count_0 = 0
    sop_files = []

    if unit:
        unit_records_qs = QualityRecord.objects.filter(unit=unit).select_related(
            'standard_item__category'
        ).prefetch_related(
            'standard_item__evidence_reqs__files'
        ).order_by('standard_item__category__order', 'standard_item__order')

        for rec in unit_records_qs:
            unit_records.append(rec)
            unit_items.append(rec.standard_item)
            total_score += rec.score
            total_possible += 10
            if rec.score == 10:
                count_10 += 1
            elif rec.score == 5:
                count_5 += 1
            else:
                count_0 += 1

        # SOP/Regulasi (kategori R) untuk unit ini
        sop_reqs = EvidenceReq.objects.filter(
            category_type='R',
            standard_item__record__unit=unit,
        ).prefetch_related('files').select_related('standard_item')
        for req in sop_reqs:
            for f in req.files.all():
                sop_files.append({
                    'file': f,
                    'ep_code': req.standard_item.code,
                    'req_title': req.title,
                })

    percentage = round((total_score / total_possible) * 100, 1) if total_possible > 0 else 0

    # Upload bukti: requirement EP unit yang belum punya file
    upload_targets = []
    if unit:
        reqs_unit = EvidenceReq.objects.filter(
            standard_item__record__unit=unit,
            category_type='D',
        ).select_related('standard_item').prefetch_related('files')
        for req in reqs_unit:
            if req.files.count() == 0:
                upload_targets.append(req)

    # Portofolio KPS nakes
    from accounts.models import NakesCredential
    credentials = NakesCredential.objects.filter(user_profile=profile).order_by('doc_type', '-created_at')
    # Auto-check expiry
    for cred in credentials:
        cred.auto_check_expiry()

    from accounts.forms import NakesCredentialForm
    credential_form = NakesCredentialForm()

    context = {
        'profile': profile,
        'unit': unit,
        'rs_profile': rs_profile,
        'unit_records': unit_records,
        'total_score': total_score,
        'total_possible': total_possible,
        'percentage': percentage,
        'count_10': count_10,
        'count_5': count_5,
        'count_0': count_0,
        'sop_files': sop_files,
        'upload_targets': upload_targets,
        'credentials': credentials,
        'credential_form': credential_form,
    }
    return render(request, 'akreditasi/portal_nakes.html', context)


@login_required
@require_POST
def upload_kredensial_nakes(request):
    """Nakes upload dokumen STR/SIP/Sertifikat pelatihan."""
    from accounts.forms import NakesCredentialForm
    profile = request.user.profile
    form = NakesCredentialForm(request.POST, request.FILES)
    if form.is_valid():
        cred = form.save(commit=False)
        cred.user_profile = profile

        file_obj = request.FILES.get('file')
        if file_obj:
            res = upload_to_supabase_storage(file_obj, f"kps/{profile.user.username}/{file_obj.name}")
            if res.get('success') and res.get('url'):
                cred.file_url = res['url']
                cred.file = None
            else:
                cred.file = file_obj

        cred.save()
        AuditLog.objects.create(
            user=request.user,
            aksi='UPLOAD',
            model_name='NakesCredential',
            object_repr=f"{cred.get_doc_type_display()} - {cred.title}",
            detail=f"Nakes mengunggah dokumen KPS: {cred.title}"
        )
        messages.success(request, f"Dokumen '{cred.title}' berhasil diunggah dan menunggu verifikasi.")
    else:
        messages.error(request, "Gagal mengunggah dokumen. Periksa form Anda.")

    return redirect('akreditasi:portal_nakes')


@login_required
@require_POST
def delete_kredensial_nakes(request, cred_id):
    """Nakes hapus dokumen kredensial miliknya sendiri."""
    from accounts.models import NakesCredential
    cred = get_object_or_404(NakesCredential, id=cred_id, user_profile=request.user.profile)
    title = cred.title
    cred.delete()
    AuditLog.objects.create(
        user=request.user,
        aksi='DELETE',
        model_name='NakesCredential',
        object_repr=title,
        detail="Nakes menghapus dokumen KPS"
    )
    messages.success(request, f"Dokumen '{title}' berhasil dihapus.")
    return redirect('akreditasi:portal_nakes')


@login_required
def rekap_kps_unit(request):
    """Rekap kelengkapan KPS staf per unit — untuk Admin RS & Kepala Unit."""
    profile = getattr(request.user, 'profile', None)
    if not profile or profile.role not in ('SUPER_ADMIN', 'ADMIN_RS', 'KEPALA_UNIT'):
        from django.http import HttpResponseForbidden
        return HttpResponseForbidden("Akses ditolak.")

    from accounts.models import UserProfile, NakesCredential

    # Kepala Unit hanya lihat unitnya, Admin lihat semua / filter
    selected_unit_id = request.GET.get('unit', 'ALL')
    units = UnitKerja.objects.all().order_by('code')

    nakes_qs = UserProfile.objects.filter(role='STAF_NAKES').select_related('unit_kerja', 'user').prefetch_related('credentials')

    if profile.role == 'KEPALA_UNIT' and profile.unit_kerja:
        nakes_qs = nakes_qs.filter(unit_kerja=profile.unit_kerja)
    elif selected_unit_id != 'ALL':
        nakes_qs = nakes_qs.filter(unit_kerja_id=selected_unit_id)

    # Auto-check expiry
    for nakes in nakes_qs:
        for cred in nakes.credentials.all():
            cred.auto_check_expiry()

    context = {
        'nakes_list': nakes_qs,
        'units': units,
        'selected_unit_id': selected_unit_id,
        'profile': profile,
    }
    return render(request, 'akreditasi/rekap_kps.html', context)


@login_required
@require_POST
def verify_kredensial(request, cred_id):
    """Admin RS / Kepala Unit verifikasi kredensial nakes."""
    profile = getattr(request.user, 'profile', None)
    if not profile or profile.role not in ('SUPER_ADMIN', 'ADMIN_RS', 'KEPALA_UNIT'):
        return JsonResponse({'success': False, 'error': 'Akses ditolak'}, status=403)

    from accounts.models import NakesCredential
    from accounts.forms import CredentialVerifyForm
    cred = get_object_or_404(NakesCredential, id=cred_id)
    form = CredentialVerifyForm(request.POST, instance=cred)
    if form.is_valid():
        form.save()
        AuditLog.objects.create(
            user=request.user,
            aksi='UPDATE',
            model_name='NakesCredential',
            object_repr=f"{cred.get_doc_type_display()} - {cred.title}",
            detail=f"Status diubah ke {cred.get_status_display()} oleh {request.user.username}"
        )
        messages.success(request, f"Status dokumen '{cred.title}' berhasil diperbarui.")
    else:
        messages.error(request, "Gagal memperbarui status.")
    return redirect('akreditasi:rekap_kps')
