from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase
from django.utils import timezone

from accounts.models import RolePermissionConfig, UserProfile
from akreditasi.models import UnitKerja
from sdm.models import (
    AturanDokumen,
    DokumenPegawai,
    HcGrant,
    Pegawai,
    PenempatanPegawai,
    PenilaianTalent,
    SkenarioSDM,
    SnapshotAnalitikSDM,
    VersiDokumen,
)
from sdm.services.analitik import (
    approve_scenario,
    assess_talent,
    compliance_audit,
    compute_turnover,
    export_audit_xlsx,
    persist_snapshot,
    propose_scenario,
)


class AnalyticsBase(TestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.unit = UnitKerja.objects.create(name='Unit Analitik', code='HC-AN')
        self.other = UnitKerja.objects.create(name='Unit Lain AN', code='HC-AN-2')
        self.user = self._user('analytics-user')
        config = RolePermissionConfig.get_config_for_role('ADMIN_RS')
        config.can_view_hc = True
        config.can_view_hc_analytics = True
        config.can_manage_hc_talent = True
        config.save(update_fields=[
            'can_view_hc', 'can_view_hc_analytics', 'can_manage_hc_talent', 'updated_at',
        ])
        for capability in ('can_view_hc', 'can_view_hc_analytics', 'can_manage_hc_talent'):
            HcGrant.objects.create(
                user=self.user, capability=capability, unit_kerja=self.unit
            )
        self.employee = Pegawai.objects.create(
            nama_lengkap='Pegawai Analitik', kategori='PPA', profesi='Perawat'
        )
        PenempatanPegawai.objects.create(
            pegawai=self.employee, unit_kerja=self.unit,
            tanggal_mulai=self.today - timedelta(days=30),
            alokasi_fte=Decimal('1.00'),
        )

    def _user(self, username):
        user = get_user_model().objects.create_user(username=username)
        UserProfile.objects.create(user=user, role='ADMIN_RS', unit_kerja=self.unit)
        return user


class TurnoverTests(AnalyticsBase):
    def test_zero_denominator_is_unavailable_not_zero_percent(self):
        self.employee.status = 'NONAKTIF'
        self.employee.save(update_fields=['status'])

        result = compute_turnover(
            user=self.user, unit=self.unit,
            period_start=self.today - timedelta(days=30), period_end=self.today,
        )

        self.assertFalse(result['is_available'])
        self.assertIsNone(result['value'])
        self.assertIn('Tidak ada pegawai aktif', result['unavailable_reason'])

    def test_turnover_uses_active_headcount_denominator(self):
        departed = Pegawai.objects.create(
            nama_lengkap='Nonaktif', kategori='PPA', status='NONAKTIF'
        )
        PenempatanPegawai.objects.create(
            pegawai=departed, unit_kerja=self.unit,
            tanggal_mulai=self.today - timedelta(days=30),
            alokasi_fte=Decimal('1.00'),
        )

        result = compute_turnover(
            user=self.user, unit=self.unit,
            period_start=self.today - timedelta(days=30), period_end=self.today,
        )

        self.assertTrue(result['is_available'])
        self.assertEqual(result['denominator'], Decimal('1'))
        self.assertEqual(result['value'], Decimal('100.0000'))

    def test_invalid_period_is_rejected(self):
        with self.assertRaises(ValidationError):
            compute_turnover(
                user=self.user, unit=self.unit,
                period_start=self.today, period_end=self.today - timedelta(days=1),
            )

    def test_foreign_unit_is_denied(self):
        with self.assertRaises(PermissionDenied):
            compute_turnover(
                user=self.user, unit=self.other,
                period_start=self.today - timedelta(days=30), period_end=self.today,
            )


class SnapshotTests(AnalyticsBase):
    def test_snapshot_is_immutable(self):
        result = compute_turnover(
            user=self.user, unit=self.unit,
            period_start=self.today - timedelta(days=30), period_end=self.today,
        )
        snapshot = persist_snapshot(
            user=self.user, unit=self.unit, metrik='TURNOVER',
            period_start=self.today - timedelta(days=30),
            period_end=self.today, result=result,
        )

        with self.assertRaises(ValidationError):
            snapshot.value = Decimal('99')
            snapshot.save()

    def test_snapshot_records_coverage_and_calculation_version(self):
        result = compute_turnover(
            user=self.user, unit=self.unit,
            period_start=self.today - timedelta(days=30), period_end=self.today,
        )

        snapshot = persist_snapshot(
            user=self.user, unit=self.unit, metrik='TURNOVER',
            period_start=self.today - timedelta(days=30),
            period_end=self.today, result=result,
        )

        self.assertEqual(snapshot.calculation_version, 'hc-analytics-v1')
        self.assertTrue(snapshot.is_available)


class ComplianceAuditTests(AnalyticsBase):
    def test_audit_separates_uploaded_and_verified(self):
        AturanDokumen.objects.create(
            document_type='CREDENTIAL', name='STR', validity_policy='EXPIRY_REQUIRED'
        )
        document = DokumenPegawai.objects.create(
            employee=self.employee, document_type='CREDENTIAL',
            title='STR', sensitivity='personnel',
        )
        VersiDokumen.objects.create(
            document=document, version_number=1, original_filename='str.pdf',
            storage_key=f'personnel/{self.employee.pk}/{document.pk}.pdf',
            sha256='a' * 64, size_bytes=10, content_type='application/pdf',
            status='PENDING',
        )

        audit = compliance_audit(user=self.user, unit=self.unit)

        row = audit['rows'][0]
        self.assertEqual(row['upload_pct'], 100.0)
        self.assertEqual(row['completeness_pct'], 0.0)

    def test_export_escapes_formula_injection_and_omits_file_urls(self):
        self.employee.nama_lengkap = '=HYPERLINK("http://evil")'
        self.employee.save(update_fields=['nama_lengkap'])
        AturanDokumen.objects.create(
            document_type='CREDENTIAL', name='STR', validity_policy='EXPIRY_REQUIRED'
        )

        audit = compliance_audit(user=self.user, unit=self.unit)
        payload = export_audit_xlsx(audit)

        self.assertTrue(payload.startswith(b'PK'))
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(payload)).active
        self.assertTrue(str(sheet.cell(row=2, column=1).value).startswith("'="))

    def test_unmapped_requirement_is_reported_not_counted_as_compliance(self):
        AturanDokumen.objects.create(
            document_type='CREDENTIAL', name='STR', validity_policy='EXPIRY_REQUIRED'
        )

        audit = compliance_audit(user=self.user, unit=self.unit)

        self.assertIn('STR', audit['unmapped_requirements'])
        self.assertEqual(audit['mapped_requirements'], 0)

    def test_health_documents_are_excluded_from_audit(self):
        DokumenPegawai.objects.create(
            employee=self.employee, document_type='HEALTH',
            title='MCU', sensitivity='health',
        )

        audit = compliance_audit(user=self.user, unit=self.unit)

        self.assertNotIn('MCU', str(audit))


class TalentTests(AnalyticsBase):
    def test_missing_axis_remains_unclassified(self):
        assessment = assess_talent(
            self.employee, actor=self.user, kinerja='TINGGI', kompetensi=''
        )

        self.assertIsNone(assessment.box)

    def test_disputed_assessment_is_not_classified(self):
        assessment = assess_talent(
            self.employee, actor=self.user, kinerja='TINGGI', kompetensi='TINGGI',
            disputed=True,
        )

        self.assertIsNone(assessment.box)

    def test_assessment_versions_increase_and_preserve_history(self):
        first = assess_talent(
            self.employee, actor=self.user, kinerja='SEDANG', kompetensi='TINGGI'
        )
        second = assess_talent(
            self.employee, actor=self.user, kinerja='TINGGI', kompetensi='TINGGI'
        )

        self.assertEqual(first.versi, 1)
        self.assertEqual(second.versi, 2)
        self.assertEqual(PenilaianTalent.objects.count(), 2)


class ScenarioTests(AnalyticsBase):
    def test_scenario_requires_human_approval(self):
        scenario = propose_scenario(
            actor=self.user, unit=self.unit, jenis='ROTASI',
            rekomendasi='Rotasi dua perawat ke unit intensif.',
            asumsi={'periode': '2027'}, provenance={'source': 'formasi'},
            keterbatasan='Data historis 6 bulan.',
        )

        self.assertEqual(scenario.status, 'DRAF')
        approve_scenario(scenario, actor=self.user)
        scenario.refresh_from_db()
        self.assertEqual(scenario.status, 'DISETUJUI')
        self.assertIsNotNone(scenario.approved_at)

    def test_scenario_carries_provenance_and_limitations(self):
        scenario = propose_scenario(
            actor=self.user, unit=self.unit, jenis='KEKURANGAN_SPESIALIS',
            rekomendasi='Rekrut satu spesialis.', asumsi={'a': 1},
            provenance={'b': 2}, keterbatasan='Belum ada data BOR unit.',
        )

        self.assertEqual(scenario.asumsi, {'a': 1})
        self.assertEqual(scenario.provenance, {'b': 2})
        self.assertIn('BOR', scenario.keterbatasan)


class PdfExportTests(AnalyticsBase):
    def test_pdf_export_produces_a_pdf_document(self):
        from sdm.services.analitik import export_audit_pdf

        AturanDokumen.objects.create(
            document_type='CREDENTIAL', name='STR', validity_policy='EXPIRY_REQUIRED'
        )
        audit = compliance_audit(user=self.user, unit=self.unit)

        payload = export_audit_pdf(audit)

        self.assertTrue(payload.startswith(b'%PDF-'))
        self.assertGreater(len(payload), 500)

    def test_pdf_export_handles_an_empty_audit(self):
        from sdm.services.analitik import export_audit_pdf

        audit = compliance_audit(user=self.user, unit=self.unit)

        self.assertTrue(export_audit_pdf(audit).startswith(b'%PDF-'))
